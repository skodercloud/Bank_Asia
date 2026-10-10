import docx

doc = docx.Document('RFP_LMS Bank Asia_v0.3_backup.docx')

print(f"Total tables: {len(doc.tables)}")

t10 = doc.tables[10]
print(f"Table 10 rows: {len(t10.rows)}")

t12 = doc.tables[12]
print(f"Table 12 rows: {len(t12.rows)}")

with open('inspect_t12.txt', 'w', encoding='utf-8') as f:
    for idx, row in enumerate(t12.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
        # dedup adjacent
        dedup = []
        for c in cells:
            if not dedup or dedup[-1] != c:
                dedup.append(c)
        f.write(f"R{idx:03d} (len={len(dedup)}): {dedup}\n")

print("Wrote inspect_t12.txt")

with open('inspect_t12_xml.txt', 'w', encoding='utf-8') as f:
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    for r in range(50):
        tr = t12._tbl.findall('w:tr', ns)[r]
        tcs = tr.findall('w:tc', ns)
        texts = [''.join(c.itertext()).strip() for c in tcs]
        f.write(f"R{r:03d} (tc_cnt={len(tcs)}): {texts}\n")

print("Wrote inspect_t12_xml.txt")
